#!/usr/bin/env python3
"""Relatório PDF (Real Audit Tech): despesas de UNIFORMES + MATERIAL DE LIMPEZA por EMPRESA.

Separa as despesas por empresa (código de cliente na coluna G), com valor final por empresa,
a parte 'SEM EMPRESA RELACIONADA' e um gráfico de pizza no final.
"""
import argparse, json, os, re
from datetime import date
import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer,
                                Table, TableStyle, HRFlowable)
from reportlab.graphics.shapes import Drawing, String
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.charts.legends import Legend

INK = colors.HexColor("#173a3b"); INK2 = colors.HexColor("#0c292a")
CORAL = colors.HexColor("#d66e54"); MINT = colors.HexColor("#9bc8b5")
PAPER = colors.HexColor("#f5f2e9"); LINE = colors.HexColor("#d8d5c9")
GREEN = colors.HexColor("#2e7d54"); RED = colors.HexColor("#c0492f")
SEM = "SEM EMPRESA RELACIONADA"
# paleta para a pizza
PIE = [colors.HexColor(c) for c in ("#173a3b", "#d66e54", "#9bc8b5", "#2e7d54", "#c07a00",
        "#5a8a8b", "#e0a48f", "#3f6f5a", "#b0563f", "#7fa9a0", "#c9b27a", "#8c4a3a")]

def brl(v):
    if v is None:
        return "—"
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-" if v < 0 else "") + s

def header_footer(canvas, doc):
    canvas.saveState()
    w, h = A4
    canvas.setFillColor(INK); canvas.rect(0, h - 16 * mm, w, 16 * mm, fill=1, stroke=0)
    canvas.setFillColor(CORAL); canvas.rect(0, h - 16 * mm, 2.6 * mm, 16 * mm, fill=1, stroke=0)
    canvas.setFillColor(CORAL); canvas.rect(18 * mm, h - 12.9 * mm, 6.6 * mm, 6.6 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.white); canvas.setFont("Helvetica-Bold", 6.6)
    canvas.drawCentredString(18 * mm + 3.3 * mm, h - 11.0 * mm, "RAT")
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(27.5 * mm, h - 9.8 * mm, "REAL AUDIT TECH")
    canvas.setFont("Helvetica", 5.8); canvas.setFillColor(MINT)
    canvas.drawString(27.7 * mm, h - 13.2 * mm, "TECNOLOGIA TRIBUTÁRIA")
    canvas.setFont("Helvetica", 8); canvas.setFillColor(MINT)
    canvas.drawRightString(w - 18 * mm, h - 10.5 * mm, f"Uniformes & Limpeza · APPA 2025 · Pág. {doc.page}")
    canvas.setStrokeColor(LINE); canvas.setLineWidth(0.5); canvas.line(18 * mm, 12 * mm, w - 18 * mm, 12 * mm)
    canvas.setFillColor(colors.HexColor("#738080")); canvas.setFont("Helvetica", 6.5)
    canvas.drawString(18 * mm, 8.5 * mm, "Uso restrito · documento gerencial. Despesas de uniformes e material de limpeza por empresa.")
    canvas.drawRightString(w - 18 * mm, 8.5 * mm, date.today().strftime("%d/%m/%Y"))
    canvas.restoreState()

REF_RE = re.compile(r"ref\.?\s*\d+", re.I)
NUM_RE = re.compile(r"\d{2,4}")

def num(v):
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "") else 0.0
    except Exception:
        return 0.0

def carrega(src, code2name):
    known = set(code2name)
    def resolve(g):
        limpo = REF_RE.sub(" ", str(g or ""))
        for tok in NUM_RE.findall(limpo):
            if tok in known:
                return tok, code2name[tok]
        return None, SEM
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    agg = {}
    for cat, sn in (("Limpeza", "limpeza"), ("Uniformes", "uniformes")):
        ws = wb[sn]; first = True
        for r in ws.iter_rows(values_only=True):
            if first:
                first = False; continue
            if not r or all(x is None for x in r):
                continue
            row = list(r) + [None] * 9
            cod, emp = resolve(row[6])
            v = num(row[7])
            d = agg.setdefault(emp, {"cod": cod, "Limpeza": 0.0, "Uniformes": 0.0, "n": 0})
            d[cat] += v; d["n"] += 1
            if cod:
                d["cod"] = cod
    wb.close()
    return agg

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=r"C:\Users\xandao\Downloads\uniformes e material de limpeza.xlsx")
    ap.add_argument("--mapa", default=r"storage\private\fiscal_auditor\appa\cruzamento_cliente.json")
    ap.add_argument("--out", default=r"storage\private\fiscal_auditor\appa\APPA_DESPESAS_UNIFORMES_LIMPEZA_2025.pdf")
    args = ap.parse_args()

    cli = json.load(open(args.mapa, encoding="utf-8"))
    code2name = {str(r["client_code"]): r["cliente"] for r in cli["rows"] if r.get("client_code")}
    agg = carrega(args.src, code2name)

    def tot(d):
        return d["Limpeza"] + d["Uniformes"]
    empresas = sorted([(k, v) for k, v in agg.items() if k != SEM], key=lambda kv: -tot(kv[1]))
    sem = agg.get(SEM)
    t_limp = sum(v["Limpeza"] for v in agg.values())
    t_unif = sum(v["Uniformes"] for v in agg.values())
    t_ger = t_limp + t_unif
    t_com = sum(tot(v) for k, v in agg.items() if k != SEM)
    t_sem = tot(sem) if sem else 0.0

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=INK, fontSize=22, leading=25, alignment=0)
    h1s = ParagraphStyle("h1s", parent=styles["Title"], textColor=CORAL, fontSize=13, leading=15, alignment=0)
    kick = ParagraphStyle("kick", parent=styles["Normal"], textColor=CORAL, fontSize=9, fontName="Helvetica-Bold", spaceAfter=8)
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.HexColor("#36595a"), fontSize=10, leading=14)
    body = ParagraphStyle("body", parent=styles["Normal"], textColor=INK2, fontSize=8.6, leading=12.5)
    sec = ParagraphStyle("sec", parent=styles["Heading2"], textColor=INK, fontSize=13, leading=16, spaceBefore=10, spaceAfter=5)
    cli_st = ParagraphStyle("cli", parent=body, fontSize=7.4, leading=8.8)

    doc = BaseDocTemplate(args.out, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                          topMargin=22 * mm, bottomMargin=16 * mm)
    doc.addPageTemplates([PageTemplate(id="m", frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height)],
                                       onPage=header_footer)])
    el = [Spacer(1, 6)]
    el.append(Paragraph("RELATÓRIO DE DESPESAS · UNIFORMES & MATERIAL DE LIMPEZA", kick))
    el.append(Paragraph("Despesas por Empresa", h1))
    el.append(Paragraph("APPA · Exercício 2025", h1s))
    el.append(HRFlowable(width="100%", thickness=1.6, color=CORAL, spaceBefore=6, spaceAfter=9))
    el.append(Paragraph(
        "Despesas de compra de uniformes e material de limpeza (valores pagos), alocadas por empresa "
        "(cliente) a partir do código de cliente registrado em cada nota. As despesas internas/matriz "
        "(sem cliente vinculado) aparecem agrupadas em <b>SEM EMPRESA RELACIONADA</b>.", sub))
    el.append(Spacer(1, 8))

    # KPIs
    kpi = [["Material de Limpeza", "Uniformes", "TOTAL GERAL", "Alocado a empresa", "Sem empresa"],
           [f"R$ {brl(t_limp)}", f"R$ {brl(t_unif)}", f"R$ {brl(t_ger)}",
            f"R$ {brl(t_com)}", f"R$ {brl(t_sem)}"]]
    kt = Table(kpi, colWidths=[36 * mm, 32 * mm, 36 * mm, 34 * mm, 30 * mm])
    kt.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), INK), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (2, 1), (2, 1), CORAL), ("TEXTCOLOR", (2, 1), (2, 1), colors.white),
        ("FONTNAME", (2, 1), (2, 1), "Helvetica-Bold"),
        ("BACKGROUND", (4, 1), (4, 1), colors.HexColor("#f3d9d2")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (1, 1), [PAPER]),
        ("BOX", (0, 0), (-1, -1), 0.5, LINE), ("INNERGRID", (0, 0), (-1, -1), 0.4, LINE),
    ]))
    el.append(kt)
    el.append(Spacer(1, 10))

    # Tabela por empresa
    el.append(Paragraph("Despesas por empresa (cliente)", sec))
    hdr = ["#", "Empresa (cliente)", "Cód.", "Limpeza", "Uniformes", "Valor final", "%"]
    data = [hdr]
    for i, (k, v) in enumerate(empresas, 1):
        t = tot(v)
        data.append([str(i), Paragraph(k, cli_st), v["cod"] or "—",
                     brl(v["Limpeza"]), brl(v["Uniformes"]), brl(t),
                     f"{t / t_ger * 100:.1f}%"])
    if sem:
        data.append(["", Paragraph("<b>SEM EMPRESA RELACIONADA</b> (interno/matriz)", cli_st), "—",
                     brl(sem["Limpeza"]), brl(sem["Uniformes"]), brl(t_sem), f"{t_sem / t_ger * 100:.1f}%"])
    data.append(["", Paragraph("<b>TOTAL GERAL</b>", cli_st), "", brl(t_limp), brl(t_unif), brl(t_ger), "100%"])

    cw = [7 * mm, 74 * mm, 12 * mm, 26 * mm, 26 * mm, 27 * mm, 14 * mm]
    t = Table(data, colWidths=cw, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), INK), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.6),
        ("ALIGN", (3, 0), (6, -1), "RIGHT"), ("ALIGN", (0, 0), (0, -1), "CENTER"), ("ALIGN", (2, 0), (2, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.6), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.6),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 1), (-1, -3), [colors.white, PAPER]),
        ("LINEABOVE", (0, -1), (-1, -1), 0.7, INK), ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#eef3ef")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]
    if sem:
        style.append(("BACKGROUND", (0, -2), (-1, -2), colors.HexColor("#f3d9d2")))
        style.append(("FONTNAME", (0, -2), (-1, -2), "Helvetica-Bold"))
    t.setStyle(TableStyle(style))
    el.append(t)
    el.append(Spacer(1, 12))

    # Gráfico de pizza — valor final por empresa (top 10 + Outras + Sem empresa)
    el.append(Paragraph("Distribuição do valor final por empresa", sec))
    TOPN = 10
    fatias = [(k, tot(v)) for k, v in empresas[:TOPN]]
    outras = sum(tot(v) for k, v in empresas[TOPN:])
    if outras > 0:
        fatias.append((f"Outras {len(empresas) - TOPN} empresas", outras))
    if t_sem > 0:
        fatias.append((SEM, t_sem))

    d = Drawing(500, 230)
    pie = Pie()
    pie.x = 8; pie.y = 15
    pie.width = 200; pie.height = 200
    pie.data = [f[1] for f in fatias]
    pie.labels = None
    pie.slices.strokeColor = colors.white; pie.slices.strokeWidth = 1
    for i in range(len(fatias)):
        pie.slices[i].fillColor = PIE[i % len(PIE)]
    pie.sideLabels = False
    d.add(pie)
    # legenda manual com valores
    leg = Legend()
    leg.x = 235; leg.y = 205
    leg.dx = 8; leg.dy = 8; leg.deltay = 15; leg.fontName = "Helvetica"; leg.fontSize = 8
    leg.boxAnchor = "nw"; leg.columnMaximum = 14; leg.alignment = "right"
    leg.colorNamePairs = [
        (PIE[i % len(PIE)], f"{(k[:34] + '…') if len(k) > 35 else k}  ·  R$ {brl(val)} ({val / t_ger * 100:.1f}%)")
        for i, (k, val) in enumerate(fatias)
    ]
    d.add(leg)
    el.append(d)
    el.append(Spacer(1, 6))
    el.append(Paragraph(
        f"Valor final total: <b>R$ {brl(t_ger)}</b> — {len(empresas)} empresas com despesa alocada "
        f"(R$ {brl(t_com)} · {t_com / t_ger * 100:.1f}%) e R$ {brl(t_sem)} sem empresa relacionada "
        f"({t_sem / t_ger * 100:.1f}%).", body))

    doc.build(el)
    print("OK ->", args.out)

if __name__ == "__main__":
    main()
